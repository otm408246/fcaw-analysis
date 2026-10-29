# -*- coding: utf-8 -*-
"""
Anade H-071 a H-075: auditoria de Zavaleta Yacila, A. M. (2013),
"Optimizacion del ancho del cordon de soldadura en acero A36 en el proceso
de soldadura FCAW, usando metodologia Taguchi". Tesis de Ingeniero
Estadistico, Universidad Nacional de Trujillo, Peru. 75 p.
Experimento en los Servicios Industriales de la Marina (SIMA).

NO ENTRA EN LA MAESTRA: junta en T (filete), ancho medido con galga, sin
diametro de alambre, sin gas, sin alimentacion. Entra como CASO DE AUDITORIA.
Todas las cifras de abajo se verificaron recalculando sobre la Tabla 07 del
propio documento (anexo, p. 62).
"""
from openpyxl import load_workbook

LIBRO = "FCAW_BASE_FINAL.xlsx"
HOJA  = "09_HALLAZGOS"

H = [
    ("H-071", "ZAVALETA2013 (tesis UNT)", "error",
     "La relacion senal-ruido con n=1 es una TAUTOLOGIA: es el ancho reetiquetado, no informacion nueva",
     "Con n=1 el criterio 'mayor es mejor' colapsa: S/R = -10*log(1/Y^2) = 20*log10(Y), una transformacion "
     "estrictamente monotona del ancho. Verificado sobre las 9 corridas de la Tabla 07: 20*log10(Y) reproduce "
     "los nueve valores publicados con desviacion maxima de 0,00005 dB, y la correlacion de rangos Y vs S/R "
     "es exactamente 1,0000. Consecuencia: las Tablas 05, 06, 07 y 08 y las Figuras 10 y 11 -la mitad del "
     "capitulo de resultados- son el mismo analisis repetido. El autor escribe (p. 53): 'notamos que el estudio "
     "arroja la misma combinacion de factores' y lo presenta como CONFIRMACION, cuando era matematicamente "
     "imposible que arrojara otra cosa. La propia tesis declara n=1 en la p. 37.",
     "HERRAMIENTA DE AUDITORIA REUTILIZABLE, del mismo tipo que la huella de imputacion: si S/R = 20*log10(Y) "
     "exacto, entonces no hubo replicas y el analisis Taguchi es decorativo. Revisar cuantos de los 27 estudios "
     "del corpus aplican S/R sin replicas. Refuerza H-004: la sofisticacion metodologica no anade precision",
     "critico", "Cap. 4 auditoria + herramienta metodologica. Caso de apertura del Cap. 1"),

    ("H-072", "ZAVALETA2013 (tesis UNT)", "error",
     "La ecuacion 'en factores reales' nunca se decodifico: el entregable es inaplicable",
     "La tesis imprime dos ecuaciones (p. 43). La segunda, declarada 'en terminos de factores reales', es la "
     "PRIMERA con las palabras sustituidas: los coeficientes son identicos y falta la sustitucion "
     "X = (Z-(Zmin-dZ))/dZ. Evaluada tal como esta impresa, para 300 A, 26 V y 900 mm/min devuelve "
     "Y = 121 144 mm de ancho, frente a los 12,02 mm realmente medidos: un factor de 10 000. "
     "La ecuacion decodificada nunca aparece en el documento.",
     "El unico objeto transferible del trabajo -la ecuacion que otro taller usaria- esta roto, y nadie lo "
     "detecto en la sustentacion. Es la version local y verificable de H-037: publicar el modelo en una forma "
     "que nadie puede aplicar. Caso de apertura ideal para el Cap. 1 porque es peruano, accesible y concreto",
     "critico", "Cap. 1 caso de apertura + Cap. 4. Refuerza H-037"),

    ("H-073", "ZAVALETA2013 (tesis UNT)", "error",
     "Todos los coeficientes publicados contradicen en SIGNO a los datos de la propia tesis",
     "Reajustando la regresion sobre las 9 corridas con codificacion centrada: corriente b1 = +1,770 "
     "(publicado -1,8267) y velocidad b3 = -0,902 (publicado +0,8333); los cuadraticos tambien invertidos. "
     "La constante publicada 9,4733 SI coincide exactamente con la media global del ancho, de modo que el "
     "ajuste es el correcto y lo alterado son los signos. Los signos publicados contradicen ademas la Tabla 01 "
     "del mismo documento (corriente 200->7,647 y 300->11,187 mm, creciente) y el texto, que dice correctamente "
     "que mas corriente da cordon mas ancho. El autor ADVIRTIO el problema y lo dejo escrito (p. 53): 'las "
     "incongruencias que existen entre la regresion de Taguchi y la regresion comun', sin resolverlo.",
     "Un revisor con los datos delante lo detecta en diez minutos; sin los datos es indetectable. Es el "
     "argumento empirico de por que la tesis propia publica la base completa y no solo los coeficientes. "
     "Sostiene la norma minima de reporte del Cap. 4",
     "alto", "Cap. 4 norma de reporte + justificacion de publicar los datos"),

    ("H-074", "ZAVALETA2013 (tesis UNT)", "contradiccion",
     "El resumen contradice a las conclusiones en el resultado principal, y el resumen es el que se indexa",
     "El Resumen (p. vi) y el Abstract (p. vii) declaran que 300 A y 1000 mm/min maximizan el ancho. El "
     "capitulo de analisis (p. 53) y las conclusiones declaran 300 A, 24 V y 900 mm/min. Los datos dan la "
     "razon a las conclusiones: medias del ancho por nivel de velocidad 900 -> 10,307 mm, 950 -> 9,610 mm, "
     "1000 -> 8,503 mm, monotona decreciente, maximo en 900. El resumen esta equivocado. "
     "ERRATA ADICIONAL: la Tabla 07 (p. 62) anota 100 mm/min en el ensayo 7; el arreglo L9 exige nivel 3, "
     "es decir 1000 mm/min.",
     "El resumen es lo unico que leen el buscador, el indexador y el siguiente tesista. Un lector que cite "
     "este trabajo propagara el valor equivocado. Suma al catalogo de erratas del corpus (H-035, H-036, H-043) "
     "un caso nuevo: la contradiccion INTERNA entre resumen y conclusiones",
     "alto", "Cap. 4 catalogo de erratas + Cap. 8 recomendaciones de reporte"),

    ("H-075", "ZAVALETA2013 vs H-034", "verificacion",
     "Confirmacion independiente del techo de precision: el residuo publicado cae POR DEBAJO del piso fisico",
     "La tesis declara n=1, sin replicas (p. 37). Su ANOVA (Tabla 02) asigna 2 GL a 'Error' con CM = 0,06173, "
     "de donde el residuo es 0,2485 mm sobre una media de 9,4733 mm, esto es 2,62%. El piso de error del ancho "
     "medido en replicas VERDADERAS de 8 estudios del corpus (H-034) es 3,84%. El residuo publicado esta 1,46 "
     "veces POR DEBAJO de la repetibilidad fisica del proceso, y sobre el se construyen el R2 = 99,5% y los "
     "valores p de 0,007 y 0,024. Sin replicas, ese 'Error' no es error experimental sino interaccion no "
     "estimada, que Taguchi asume nula por postulado.",
     "CONFIRMACION INDEPENDIENTE DE H-004 Y H-034 con un estudio ajeno al corpus de modelado: un modelo cuyo "
     "error es menor que el ruido del proceso esta memorizando, no explicando. Es la demostracion mas limpia "
     "disponible porque el propio documento publica el CM del error y declara la ausencia de replicas",
     "critico", "Cap. 4 piso de error + Cap. 5 criterio de parada. Confirma H-004 y H-034"),
]

wb = load_workbook(LIBRO)
ws = wb[HOJA]
ya = {r[0].value for r in ws.iter_rows(min_col=1, max_col=1)}
for h in H:
    if h[0] in ya:
        raise SystemExit(f"{h[0]} ya existe; nada que hacer")
    ws.append(list(h))
wb.save(LIBRO)
print(f"{HOJA}: {ws.max_row - 1} hallazgos")
