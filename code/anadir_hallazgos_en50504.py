# -*- coding: utf-8 -*-
"""
Anade H-064 a H-070 a la hoja 09_HALLAZGOS: revision de BS EN 50504:2008
"Validation of arc welding equipment" (CENELEC TC 26A, aprobada 2008-06-01).

Es el documento al que ISO 17662:2016 Tabla 11 remite para validar
amperimetros y voltimetros. Se obtuvo escaneado; se leyo renderizando
las 29 paginas a imagen.
"""
from openpyxl import load_workbook

LIBRO = "FCAW_BASE_FINAL.xlsx"
HOJA  = "09_HALLAZGOS"

H = [
    ("H-064", "EN 50504:2008 Foreword + ISO 17662:2016 Tabla 11", "contradiccion",
     "La obligacion de validar es imperativa pero su contenido numerico es, en la fuente, una recomendacion",
     "ISO 17662:2016 Tabla 11 exige ('SHALL BE VALIDATED') validar amperimetros y voltimetros y remite a EN 50504. "
     "Pero EN 50504:2008, Foreword, declara literalmente: 'As a code of practice, this European Standard takes the form "
     "of guidance and recommendations. It should not be quoted as if it were a specification and particular care should "
     "be taken to ensure that claims of compliance are not misleading.' Ademas los Anexos A, B, C, E y F -donde estan "
     "las tolerancias de grado precision y TODO lo relativo al alimentador de alambre- son INFORMATIVOS.",
     "HUECO NORMATIVO EN LA CADENA: la exigencia es mandatoria, el numero al que remite no lo es. Un taller puede "
     "declarar cumplimiento de ISO 3834-2 sin que exista un umbral exigible contra el cual medirlo. Lo que la tesis "
     "aporta -un criterio de verificacion reproducible por probeta- ocupa exactamente ese vacio",
     "critico", "NUCLEO CAP. 6 + Cap. 8 aportacion normativa"),

    ("H-065", "EN 50504:2008 Tabla 1 y Tabla A.1 vs AWS D1.1:2020 Tabla 6.5", "contradiccion",
     "El grado estandar de EN 50504 NO alcanza el umbral de voltaje de AWS D1.1; hace falta grado precision",
     "EN 50504:2008 Tabla 1 (grado estandar): corriente y voltaje +/-10% del valor verdadero entre el 100% y el 25% del "
     "ajuste maximo, y +/-2,5% del ajuste maximo por debajo del 25%. Tabla A.1 (grado precision, Anexo A informativo): "
     "corriente +/-2,5% (100-40%) y +/-1% (bajo 40%); voltaje +/-5% (100-40%) y +/-2% (bajo 40%). "
     "AWS D1.1/D1.1M:2020 Tabla 6.5 para FCAW: voltaje >7% obliga a recalificar (item 16); corriente >10% (item 12). "
     "CONFRONTACION: 10% > 7%. Una maquina PLENAMENTE CONFORME a EN 50504 grado estandar puede estar fuera de la "
     "variable esencial de voltaje de AWS D1.1. En corriente ambos coinciden en 10%, justo en el limite.",
     "RESULTADO ACCIONABLE: para sostener un WPS bajo AWS D1.1 en FCAW no basta la validacion de grado estandar; "
     "se requiere GRADO PRECISION (+/-5% en voltaje). Se especifica asi en el sistema de la tesis",
     "critico", "Cap. 6 especificacion de la calibracion + Cap. 8"),

    ("H-066", "EN 50504:2008 Anexo B.1 vs AWS D1.1:2020 Tabla 6.5 item 17", "contradiccion",
     "La norma declara INNECESARIO validar la velocidad de alambre, que es la variable troncal del proceso",
     "EN 50504:2008, Anexo B.1: 'It is recognized that in MIG/MAG and allied self-adjusting arc systems in which the "
     "wire feed rate is linked to a function of the welding power, AN ABSOLUTE CALIBRATION OF WIRE FEED RATE IS NOT "
     "NECESSARY.' Y B.2: 'in MIG welding the wire feed speed may be adjusted to give a specific welding current with the "
     "speed control scale in arbitrary units. IT IS NOT NECESSARY TO VALIDATE THE SPEED CONTROL on such a wire feeder.' "
     "El Anexo B es ademas INFORMATIVO. En sentido contrario, AWS D1.1:2020 Tabla 6.5 item 17 hace de la velocidad de "
     "alambre una VARIABLE ESENCIAL con umbral >10%. Tabla B.1, cuando si se valida: estandar +/-10%, precision +/-2,5%.",
     "CAUSA RAIZ NORMATIVA del problema central del corpus: el 27% de los registros tiene alimentacion imputada y varios "
     "autores la retrocalculan desde la corriente (H-045). No es descuido de los autores: NINGUNA NORMA LES OBLIGA A "
     "MEDIRLA. La variable troncal del modelo (WFS/TS) es la unica que el sistema normativo declara opcional. "
     "Es el hueco mas fuerte del bloque y el que justifica medirla por probeta",
     "critico", "NUCLEO CAP. 6 + Cap. 5 justificacion + Cap. 8 aportacion. Eleva H-045 y H-047"),

    ("H-067", "EN 50504:2008 clausula 8.4", "verificacion",
     "H-044 queda respaldado literalmente por norma: la corriente es respuesta, no consigna",
     "EN 50504:2008, 8.4 'Metal inert/active gas and flux cored arc welding': 'MIG/MAG and FCW power sources do not "
     "usually have current controls. The output current of such a welding power source is generally controlled by the "
     "wire feed speed setting which is set on the wire feeder control panel. THE CURRENT VARIES AUTOMATICALLY to "
     "regulate the burn off rate of the consumable electrode.' Coincide punto por punto con H-044, que hasta ahora se "
     "sostenia por inferencia a partir del corpus.",
     "H-044 asciende de inferencia propia a cita normativa. La especificacion del modelo con WFS/TS como predictor "
     "troncal en lugar de la corriente deja de ser una decision estadistica y pasa a ser consecuencia del proceso "
     "tal como lo describe la norma del equipo",
     "critico", "Cap. 5 justificacion del modelo + Cap. 2 marco. Consolida H-044"),

    ("H-068", "EN 50504:2008 clausula 9.5.1.9 vs ISO 17662:2016 clausula 4.2", "hallazgo",
     "La cobertura de rango SI esta normada -cinco puntos- pero en el unico documento que no obliga",
     "EN 50504:2008, 9.5.1.9 y parrafo siguiente: 'Make the measurements at the minimum setting, the maximum setting and "
     "THREE OTHER POINTS nominally equally spaced between minimum and maximum, over the full or partial range.' Es decir "
     "CINCO puntos. NOTA del mismo apartado: si se valida un rango parcial, el minimo se aborda desde abajo y el maximo "
     "desde arriba, y se promedian ambas lecturas en cada punto -control de histeresis. "
     "Pero: EN 50504 es codigo de practica (H-064); ISO 17662:2016 clausula 4.2 obliga a validar y solo fija la "
     "PERIODICIDAD anual; AWS D1.1 fija umbrales de variable esencial pero ningun metodo de validacion.",
     "CORRIGE Y AFINA H-058. El hueco no es que nadie haya pensado en la cobertura del rango: es que NINGUN documento "
     "de la cadena la hace EXIGIBLE. La propuesta de la tesis se reformula con mas fuerza: no inventar el requisito "
     "de multipunto, sino hacerlo verificable por probeta y trazable al WPS. H-031 (+13,7% a 170 A y -15,1% a 230 A) "
     "es la demostracion empirica de por que cinco puntos importan",
     "critico", "Cap. 6 diseno de la calibracion + Cap. 8. REEMPLAZA la redaccion de H-058"),

    ("H-069", "EN 50504:2008 clausula 8.4 y Anexo F vs H-031 (Solovyov2023)", "hallazgo",
     "La norma valida el voltaje en vacio, condicion en la que el proceso nunca opera; y su regla practica no reproduce lo medido",
     "EN 50504:2008, 8.4: 'The voltage control sets the NO-LOAD voltage, and the voltage will fall slightly as current is "
     "drawn (TYPICALLY 3 V TO 5 V PER 100 A)' y 'The voltage control if marked in absolute units SHOULD BE VALIDATED AT "
     "NO-LOAD VOLTAGES.' Anexo F: el voltimetro del equipo mide entre los bornes de la fuente (puntos 1-7) mientras el "
     "voltaje de arco real esta entre 4-5, de modo que 'the meter reading will therefore tend to be HIGHER than the "
     "actual arc voltage'. PRUEBA CUANTITATIVA contra Solovyov2023 (corriente medida 160 a 315 A): aplicando la regla de "
     "3-5 V/100 A, a 160 A la caida seria 4,8-8,0 V, esto es -20% a -33% sobre una consigna de 24 V, y lo observado es "
     "-19,2%, apenas por debajo del extremo inferior; pero a 315 A la caida seria 9,5-15,8 V, esto es -30% a -49% sobre "
     "32 V, y lo observado es solo -5,6%.",
     "DOBLE RESULTADO. (a) La norma valida el voltaje justo donde el proceso no trabaja: sin arco. (b) Su propio valor "
     "tipico no reproduce la maquina real -acierta en el extremo bajo y falla por un factor de 5 a 9 en el alto-. "
     "Confirma el mecanismo fisico de H-031 y a la vez demuestra que la regla practica no sustituye a la medicion. "
     "La desviacion hay que medirla en la maquina propia, que es precisamente lo que hace la probeta",
     "critico", "Cap. 6 + Cap. 8. Da mecanismo fisico a H-031"),

    ("H-070", "EN 50504:2008 clausulas 6, 9.3, 9.5.2, 10 y Anexo B.3", "oportunidad",
     "El procedimiento completo del capitulo 6 queda especificado y citable, con equipo que ya se tiene",
     "INSTRUMENTOS, 9.3 c): deben ser 'at least twice and preferably five times more accurate than the accuracy required "
     "for the validation grade'. Para grado estandar (+/-10%) implica instrumento de <=5% (2x) o <=2% (5x); para grado "
     "precision en voltaje (+/-5%) implica <=2,5% o <=1%. "
     "ALIMENTACION, Anexo B.3 metodo a): cronometrar el tiempo en que se entregan aproximadamente 1 m de alambre por la "
     "boquilla y medir la longitud con regla de acero a 1 mm. NO HACE FALTA EQUIPO ADICIONAL. "
     "SECUENCIA: 9.5.2.1 encender 5 min antes; 9.5.2.2 estabilizar 10 s por punto en la fuente; B.3 estabilizar 5 s, "
     "barrido ascendente y luego descendente, y registrar las dos lecturas y su media. "
     "CARGA, 9.5.1.4: carga resistiva, nunca cortocircuito en fuente de voltaje constante. "
     "PERIODICIDAD, clausula 6: anual, con retest a los 3 meses tras el ensayo de consistencia inicial; Anexo A NOTA: "
     "6 meses para grado precision. VALIDADORES, clausula 7: no exige certificacion de tercero, el usuario puede "
     "hacerlo con equipo trazable. SALIDA, clausula 10: etiqueta y certificado de validacion.",
     "El protocolo experimental del Cap. 6 se redacta citando clausula por clausula en lugar de proponerlo ex novo. "
     "Confirma que el equipo declarado (amperimetro, voltimetro, vernier, cronometro, regla) basta, y anade solo el "
     "requisito de EXACTITUD de los instrumentos y una carga resistiva",
     "alto", "Cap. 6 protocolo experimental completo"),
]

wb = load_workbook(LIBRO)
ws = wb[HOJA]
ya = {r[0].value for r in ws.iter_rows(min_col=1, max_col=1)}
for h in H:
    if h[0] in ya:
        raise SystemExit(f"{h[0]} ya existe; nada que hacer")
    ws.append(list(h))
wb.save(LIBRO)
print(f"{HOJA}: {ws.max_row - 1} hallazgos")
