# -*- coding: utf-8 -*-
"""
Anade H-076 a H-080: auditoria de Sucuitana Munoz, R. F. y Velasquez Zavala, R. A.
(2020), "Efecto de la aplicacion de soldadura FCAW en la productividad del soldador
area de mantenimiento de Austral Group SAA, Coishco - 2020". Tesis de Ingeniero
Industrial, Universidad Cesar Vallejo, Chimbote, Peru. 91 p.
Linea de investigacion: Gestion Empresarial y Productiva.

OJO: el archivo se llama "699532000-PROCEDIMINETOS-DE-SOLDADURA-FCAW.pdf".
El nombre no corresponde al contenido.

NO ENTRA EN LA MAESTRA: no mide ninguna propiedad del cordon. Entra como
CASO DE NEGOCIO (cuantifica el valor economico del FCAW) y como HUECO
(cambia el proceso sin verificar la calidad de la soldadura resultante).
Cifras recalculadas sobre sus Tablas 5, 11, 14 y 15.
"""
from openpyxl import load_workbook

LIBRO = "FCAW_BASE_FINAL.xlsx"
HOJA  = "09_HALLAZGOS"

H = [
    ("H-076", "AUSTRAL2020 (tesis UCV)", "error",
     "La prueba estadistica publicada contradice al diseno declarado, aunque la conclusion sobrevive",
     "El documento declara (p. vii y p. 22) un diseno cuasiexperimental con GRUPO DE CONTROL y GRUPO "
     "EXPERIMENTAL de 3 soldadores CADA UNO, es decir dos grupos independientes de personas distintas. "
     "Pero la Tabla 15 publica una prueba t PAREADA: 'SMAW - FCAW', media 11,08200, desviacion 1,67660, "
     "error 0,43290, t = 25,600, gl = 14. Recalculado sobre sus Tablas 5 y 11: la t pareada da exactamente "
     "25,600 con gl 14, de modo que el emparejamiento esta confirmado. El emparejamiento exige que la "
     "observacion i de un metodo y la i del otro provengan de la MISMA unidad, lo que es imposible entre "
     "dos grupos de soldadores distintos. La t de dos muestras independientes, que es la que su diseno exige, "
     "da t = 23,519 con gl = 28. Ademas las varianzas difieren 12,3 veces (1,76^2 frente a 0,50^2), lo que "
     "viola la homocedasticidad; la correccion de Welch da igualmente t = 23,519.",
     "DEFECTO DE REPORTE, NO DE RESULTADO: con la prueba correcta la diferencia sigue siendo significativa "
     "muy por debajo de 0,001. Sirve como ejemplo de que una conclusion puede ser correcta y su justificacion "
     "estadistica no serlo, y de por que la tesis propia publica los datos que permiten rehacer la prueba",
     "medio", "Cap. 4 auditoria + Cap. 2 antecedentes nacionales"),

    ("H-077", "AUSTRAL2020 vs el problema del ICC", "hallazgo",
     "Quince observaciones de tres soldadores analizadas como quince unidades independientes: el mismo "
     "problema de anidamiento que el corpus, en el dominio de la productividad",
     "El analisis usa gl = 14, lo que supone 15 unidades independientes. Las unidades independientes reales "
     "son 3 soldadores por grupo: las 15 observaciones estan ANIDADAS en ellos. Ademas, si los grupos son de "
     "personas distintas -como declara el documento- la comparacion SMAW frente a FCAW queda completamente "
     "CONFUNDIDA con la comparacion entre los dos grupos de soldadores: no hay forma de separar 'el FCAW es "
     "mas rapido' de 'estos tres soldadores son mas rapidos que aquellos tres'. Es estructuralmente identico "
     "al problema interlaboratorio del corpus, donde el efecto del metodo y el efecto del laboratorio no son "
     "separables sin solapamiento de condiciones.",
     "VALOR PARA EL CAP. 2: demuestra que el problema de anidamiento y confusion que la tesis propia resuelve "
     "con el modelo jerarquico NO es especifico de la geometria del cordon, sino que reaparece en el estudio "
     "de tiempos, en gestion, con la misma estructura. Es el puente natural entre la parte tecnica y la LGAC-3",
     "alto", "Cap. 2 marco + Cap. 5 justificacion del modelo jerarquico"),

    ("H-078", "AUSTRAL2020 (tesis UCV)", "error",
     "La linea base se MIDIO y la mejora se AFIRMO: el 0% de improductividad es un conteo de simbolos",
     "Para el SMAW se realizo una medicion empirica del tiempo improductivo (Figura 7): 38,54% +/- 6,29 del "
     "tiempo total de operacion, equivalente a 5,98 min +/- 1,52, calculado comparando tiempo de arco contra "
     "tiempo total en 15 trabajos. Para el FCAW esa medicion NO se repitio: el 0% de actividades improductivas "
     "(p. 38) proviene de contar las cajas del cursograma analitico, donde desaparecen las 3 demoras por cambio "
     "de electrodo. El FCAW conserva cambios de carrete, cambio de punta de contacto, cambio de cilindro de gas "
     "y -por ser alambre tubular- remocion de escoria. Un 0% exacto de tiempo improductivo no es plausible.",
     "ASIMETRIA METODOLOGICA: el diagnostico se midio, la mejora se dedujo del diagrama. El resultado principal "
     "del trabajo -la reduccion de 20,66 a 5,76 min de tiempo estandar- SI esta medido y sobrevive; lo que no "
     "esta medido es la afirmacion de improductividad nula",
     "medio", "Cap. 4 auditoria + Cap. 2 antecedentes"),

    ("H-079", "AUSTRAL2020 Figura 5 vs AWS D1.1:2020 Tabla 6.5", "verificacion",
     "Confirmacion empirica de H-051 con medicion de taller real: la soldadura manual cruza el umbral de "
     "variable esencial de velocidad de avance",
     "La Figura 5 reporta el avance manual medido en 15 trabajos de mantenimiento: fluctua entre 2,10 y "
     "3,89 cm/min con media 2,79 +/- 0,59 cm/min. Recalculado: CV = 21,1%, y los extremos se apartan de la "
     "media -24,7% y +39,4%. AWS D1.1/D1.1M:2020 Tabla 6.5 admite hasta 25% de variacion de velocidad de "
     "avance sin recalificar el procedimiento. El extremo superior observado LO CRUZA. "
     "Consistencia interna verificada: con 38,54% de tiempo improductivo sobre 15,36 min, el tiempo de arco "
     "es 9,44 min, que a 2,79 cm/min da 26,3 cm de cordon, magnitud coherente con una probeta de ensayo.",
     "Hasta ahora H-051 se apoyaba en la cifra bibliografica de +/-30% de variacion a pulso. Este documento "
     "aporta MEDICION PROPIA EN ASTILLERO PERUANO y publicada. Justifica normativamente el avance mecanizado "
     "del Cap. 6: a mano no se puede garantizar el cumplimiento de la propia variable esencial",
     "alto", "Cap. 6 diseno experimental. Confirma H-051 con dato empirico"),

    ("H-080", "AUSTRAL2020 (tesis UCV)", "oportunidad",
     "Cambiaron el proceso de soldadura de un astillero y lo metieron al sistema de calidad SIN medir una "
     "sola propiedad de la soldadura: ese es exactamente el hueco de esta tesis",
     "El trabajo sustituye SMAW por FCAW-G en el area de mantenimiento de flota de Austral Group, reduce el "
     "tiempo estandar de 20,66 a 5,76 min por probeta (-72,13%, ahorro de 14,90 min) y declara (p. 39) que "
     "'se continuo con su inclusion en el sistema de gestion de calidad de la empresa'. En las 91 paginas NO "
     "aparece: ancho, refuerzo ni penetracion del cordon; ningun ensayo macrografico; ningun WPS; ninguna "
     "calificacion segun AWS D1.1; ninguna verificacion de que el equipo entregue los parametros de consigna. "
     "La unica variable de respuesta es el TIEMPO. Se cambio el proceso y se formalizo en el sistema de "
     "calidad sobre evidencia exclusivamente de productividad.",
     "DOBLE USO PARA LA TESIS. (a) CASO DE NEGOCIO: cuantifica en una empresa real lo que vale adoptar FCAW "
     "-72% de tiempo-, que es el argumento economico que la LGAC-3 exige y que la tesis propia no tenia "
     "medido. (b) HUECO: nadie verifico que las soldaduras resultantes cumplan el codigo. Justo eso es lo que "
     "el sistema de soporte a la decision aporta. Es el mejor antecedente nacional disponible: mismo pais, "
     "mismo proceso, misma industria naval, y termina donde esta tesis empieza",
     "critico", "Cap. 1 justificacion + Cap. 2 antecedentes + Cap. 8 caso de negocio"),
]

wb = load_workbook(LIBRO)
ws = wb[HOJA]
ya = {r[0].value for r in ws.iter_rows(min_col=1, max_col=1)}
for h in H:
    if h[0] in ya:
        raise SystemExit(f"{h[0]} ya existe; nada que hacer")
    ws.append(list(h))
wb.save(LIBRO)
print(f"{HOJA}: {ws.max_row - 1} hallazgos")
